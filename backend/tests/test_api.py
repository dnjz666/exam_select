"""M3 后端 API 集成测试（AGENTS.md §7 / §10 M3 验收命令）。

覆盖：每个端点 + 三条契约铁律
1. ``recommend`` 每个 item 的 ``evidence`` 非空且带 ``source_url``；
2. ``probability is None`` ⇔ ``confidence == NO_DATA`` 且 ``reasons`` 说明原因；
3. 响应中不出现无来源的数字。

前置条件：数据库已由 ``scripts/seed.py`` 播种（M1 验收命令），否则本文件会**明确失败**
并提示运行命令——不静默跳过（"不许说应该通过了"）。
"""

from __future__ import annotations

import json
import re
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.core.rules import PROVINCES
from app.db.session import SessionLocal
from app.main import app
from app.services import chat_service

API = "/api/v1"
REPO_ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="session")
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture(scope="session", autouse=True)
def require_seeded_database() -> None:
    """M3 测试依赖已播种的数据；未播种时明确报错并给出命令。"""
    from sqlalchemy import func, select

    from app.db import models as db

    with SessionLocal() as session:
        units = session.execute(select(func.count()).select_from(db.AdmissionUnitRow)).scalar_one()
    if units == 0:
        pytest.fail(
            "数据库为空，API 集成测试无法运行。请先执行："
            "backend\\.venv\\Scripts\\python.exe scripts\\seed.py --reset"
        )


@pytest.fixture()
def student_id(client: TestClient) -> str:
    """每个测试用独立档案（避免相互污染）。"""
    payload = {
        "province": "zhejiang",
        "year": 2026,
        "subjects": ["物理", "化学", "生物"],
        "total_score": 640,
        "gender": "男",
        "single_subject_scores": {"英语": 125},
    }
    response = client.post(f"{API}/students", json=payload)
    assert response.status_code == 201, response.text
    return response.json()["data"]["id"]


# ---------------------------------------------------------------------------
# 健康 / OpenAPI
# ---------------------------------------------------------------------------
def test_health(client: TestClient) -> None:
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_openapi_schema_is_valid_and_complete(client: TestClient) -> None:
    response = client.get("/openapi.json")
    assert response.status_code == 200
    schema = response.json()
    paths = schema["paths"]
    for expected in (
        "/api/v1/meta/provinces",
        "/api/v1/meta/provinces/{province}/rule",
        "/api/v1/meta/provinces/{province}/subject-coverage",
        "/api/v1/meta/tiers",
        "/api/v1/students",
        "/api/v1/students/{student_id}/resolve-rank",
        "/api/v1/colleges/search",
        "/api/v1/majors/search",
        "/api/v1/units/{unit_id}/history",
        "/api/v1/recommend",
        "/api/v1/plans/generate",
        "/api/v1/plans/{plan_id}",
        "/api/v1/plans/{plan_id}/items",
        "/api/v1/plans/{plan_id}/validate",
        "/api/v1/plans/{plan_id}/export",
        "/api/v1/risk/scan",
        "/api/v1/chat",
        "/api/v1/chat/{session_id}/history",
        "/api/v1/backtest/report",
    ):
        assert expected in paths, expected


def test_openapi_payloads_are_typed_not_free_form(client: TestClient) -> None:
    """★ M4 前置条件：响应体必须**有类型**，否则前端"类型从 OpenAPI 生成"形同虚设。

    AGENTS.md §4.3 硬性规则 2 禁止前端手写第二份类型；若响应体是裸 ``dict``，
    生成的 TS 类型就是 ``{[key: string]: unknown}``，契约约束会整个失效。
    """
    schema = client.get("/openapi.json").json()
    components = schema["components"]["schemas"]
    for name in (
        "RecommendPayload",
        "RecommendItem",
        "RecommendStats",
        "PlanPayload",
        "PlanStats",
        "ProvinceMeta",
        "BatchMeta",
        "SubjectPoolMeta",
        "SubjectCoveragePayload",
        "StudentPayload",
        "ResolveRankPayload",
        "TiersPayload",
        "UnitHistoryPayload",
        "RiskScanPayload",
        "ChatHistoryPayload",
    ):
        assert name in components, f"OpenAPI 缺少响应模型 {name}"
        assert components[name].get("properties"), f"{name} 没有声明任何字段"

    item = components["RecommendItem"]["properties"]
    for key in (
        "unit",
        "college",
        "major",
        "probability",
        "probability_interval",
        "tier",
        "confidence",
        "utility",
        "score_breakdown",
        "predicted_min_rank",
        "sigma",
        "evidence",
        "adjustments",
        "reasons",
        "warnings",
    ):
        assert key in item, f"RecommendItem 缺字段 {key}"

    recommend_200 = schema["paths"]["/api/v1/recommend"]["post"]["responses"]["200"]
    assert "RecommendPayload" in json.dumps(recommend_200), "推荐端点未引用强类型载荷"


# ---------------------------------------------------------------------------
# 元数据
# ---------------------------------------------------------------------------
def test_meta_provinces_carries_sources_and_red_line(client: TestClient) -> None:
    body = client.get(f"{API}/meta/provinces").json()
    assert len(body["data"]) == len(PROVINCES)
    for province in body["data"]:
        assert province["batches"], province["province"]
        for batch in province["batches"]:
            assert batch["source_url"], batch["batch_code"]
            assert batch["source_quote"], batch["batch_code"]
            assert batch["verified_status"] in {"PRIMARY", "PRIMARY_GOV", "SECONDARY", "UNVERIFIED"}
    red_line = [item["province"] for item in body["data"] if item["requires_banner"]]
    assert set(red_line) == {"tianjin", "hainan"}
    assert any("规则待核实" in warning for warning in body["warnings"])
    assert body["evidence"], "元数据响应必须带来源证据"


def test_meta_province_rule_and_unknown_province(client: TestClient) -> None:
    body = client.get(f"{API}/meta/provinces/zhejiang/rule").json()
    assert body["data"]["main_batch_code"] == "zhejiang.public.seg1"
    assert body["data"]["source_problems"] == []
    missing = client.get(f"{API}/meta/provinces/jiangsu/rule")
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "NOT_FOUND"


def test_meta_tiers_documents_safety_gate(client: TestClient) -> None:
    body = client.get(f"{API}/meta/tiers").json()
    data = body["data"]
    assert {tier["tier"] for tier in data["tiers"]} >= {"CHONG", "WEN", "BAO", "DIAN"}
    assert data["safety_gate"]["affects"] == ["BAO", "DIAN"]
    assert data["safety_gate"]["warning_code"] == "SAFETY_MARGIN_NOT_MET"
    assert "仅供参考" in data["disclaimer"]


def test_meta_provinces_carries_subject_pool_with_source(client: TestClient) -> None:
    """AGENTS.md §8.1 Step 2：省份必须给出**带来源**的选考科目池。"""
    body = client.get(f"{API}/meta/provinces").json()
    by_province = {item["province"]: item for item in body["data"]}
    for province, item in by_province.items():
        pool = item["subject_pool"]
        assert pool is not None, province
        assert pool["origin"] == "RULE", province
        assert pool["source_url"] and pool["source_quote"], province
        assert pool["verified_year"] is not None, province
        assert pool["choose"] == 3 and len(pool["subjects"]) >= 3, province
        assert item["current_year"], province
    assert by_province["zhejiang"]["subject_pool"]["mode"] == "7选3"
    assert "技术" in by_province["zhejiang"]["subject_pool"]["subjects"]
    for province in ("shanghai", "beijing", "shandong", "tianjin", "hainan"):
        assert by_province[province]["subject_pool"]["mode"] == "6选3", province
        assert "技术" not in by_province[province]["subject_pool"]["subjects"], province
    # 天津科目池为 PRIMARY-GOV：必须被显式提示，不得当成 PRIMARY
    assert by_province["tianjin"]["subject_pool"]["requires_caution"] is True
    assert any("tianjin" in warning for warning in body["warnings"])
    assert any(entry["what"] == "subject_pool" for entry in body["evidence"])


def test_subject_coverage_is_data_backed(client: TestClient) -> None:
    """§8.1 Step 2 的"可报专业覆盖率"必须来自真实招生计划统计。"""
    body = client.get(
        f"{API}/meta/provinces/zhejiang/subject-coverage",
        params={"subjects": "物理,化学,生物"},
    ).json()
    data = body["data"]
    assert data["total_units"] > 0
    assert 0.0 < data["coverage"] <= 1.0
    assert data["matched_units"] <= data["total_units"]
    assert data["batch_code"] == "zhejiang.public.seg1"
    assert data["source_url"], "覆盖率必须能追到规则来源"
    assert data["subject_pool"]["subjects"]
    assert all(entry.get("source_url") for entry in body["evidence"])

    # 换一个明显更窄的组合，覆盖率不得上升（物理必选单位很多）
    narrow = client.get(
        f"{API}/meta/provinces/zhejiang/subject-coverage",
        params={"subjects": "思想政治,历史,地理"},
    ).json()["data"]
    assert narrow["coverage"] < data["coverage"]

    missing = client.get(f"{API}/meta/provinces/jiangsu/subject-coverage")
    assert missing.status_code == 404



# ---------------------------------------------------------------------------
# 考生档案
# ---------------------------------------------------------------------------
def test_create_student_draft_reports_missing_fields(client: TestClient) -> None:
    response = client.post(f"{API}/students", json={"province": "shanghai", "year": 2026})
    assert response.status_code == 201
    data = response.json()["data"]
    assert "subjects" in data["missing_fields"]
    assert "total_score" in data["missing_fields"]
    assert response.json()["warnings"], "缺字段必须给出提示"


def test_get_and_patch_student(client: TestClient, student_id: str) -> None:
    fetched = client.get(f"{API}/students/{student_id}")
    assert fetched.status_code == 200
    assert fetched.json()["data"]["missing_fields"] == []

    patched = client.patch(
        f"{API}/students/{student_id}",
        json={"subjects": ["物理", "化学", "地理"], "single_subject_scores": {"英语": 130}},
    )
    assert patched.status_code == 200
    assert patched.json()["data"]["subjects"] == ["物理", "化学", "地理"]

    assert client.get(f"{API}/students/not-exist").status_code == 404


def test_resolve_rank_traces_source(client: TestClient, student_id: str) -> None:
    response = client.post(f"{API}/students/{student_id}/resolve-rank")
    assert response.status_code == 200
    body = response.json()
    data = body["data"]
    assert data["rank"] > 0
    assert 0 < data["percentile"] <= 1
    assert data["source_url"], "位次必须可追溯来源"
    assert body["evidence"] and body["evidence"][0]["source_url"] == data["source_url"]
    assert data["equivalent_scores"], "应给出其他年份的等效分（供家长理解）"


def test_recommend_on_incomplete_profile_is_409(client: TestClient) -> None:
    draft = client.post(f"{API}/students", json={"province": "zhejiang", "year": 2026}).json()
    response = client.post(
        f"{API}/recommend", json={"student_id": draft["data"]["id"], "limit": 3}
    )
    assert response.status_code == 409
    error = response.json()["error"]
    assert error["code"] == "PROFILE_INCOMPLETE"
    assert "subjects" in error["details"]["missing_fields"]


# ---------------------------------------------------------------------------
# 数据查询
# ---------------------------------------------------------------------------
def test_colleges_and_majors_search(client: TestClient) -> None:
    colleges = client.get(f"{API}/colleges/search", params={"q": "浙江", "limit": 5}).json()
    assert colleges["data"], "应能检索到院校"
    assert all(item["source_url"] for item in colleges["data"])

    majors = client.get(f"{API}/majors/search", params={"q": "计算机", "limit": 5}).json()
    assert majors["data"]
    assert any("计算机" in item["name"] for item in majors["data"])
    assert all(item["source_url"] for item in majors["data"])

    empty = client.get(f"{API}/colleges/search", params={"q": "不存在的院校名"}).json()
    assert empty["data"] == [] and empty["warnings"]


def test_unit_history_endpoint(client: TestClient) -> None:
    with SessionLocal() as session:
        from sqlalchemy import select

        from app.db import models as db

        unit_id = session.execute(select(db.AdmissionUnitRow.unit_id).limit(1)).scalar_one()

    body = client.get(f"{API}/units/{unit_id}/history", params={"years": 3}).json()
    records = body["data"]["records"]
    assert records, "已有单位应能查到历史"
    assert len(records) <= 3
    assert all(record["source_url"] for record in records)
    assert all(record["min_rank"] is None or record["min_rank"] > 0 for record in records)
    assert body["evidence"]

    assert client.get(f"{API}/units/zhejiang-9999-NA-999999/history").status_code == 404


# ---------------------------------------------------------------------------
# 推荐（含契约铁律）
# ---------------------------------------------------------------------------
def test_recommend_contract_rules(client: TestClient, student_id: str) -> None:
    response = client.post(
        f"{API}/recommend",
        json={
            "student_id": student_id,
            "limit": 12,
            "filters": {"regions": ["zhejiang"], "majors": ["工学"]},
            "weights": {"major": 0.5, "city": 0.2},
        },
    )
    assert response.status_code == 200
    body = response.json()
    items = body["data"]["items"]
    assert items, "应能给出推荐项"
    assert len(items) <= 12

    for item in items:
        # 契约铁律 1：evidence 非空且每条带 source_url
        assert item["evidence"], item["unit"]["unit_id"]
        for entry in item["evidence"]:
            assert entry["source_url"], entry["unit"]["unit_id"]
        # 契约铁律 2：概率为 None 的项不得出现在结果里
        assert item["probability"] is not None
        assert item["confidence"] != "NO_DATA"
        # §8：概率必须以区间形式可用
        low, high = item["probability_interval"]
        assert 0.0 <= low <= high <= 1.0
        assert item["tier"] in {"CHONG", "WEN", "BAO", "DIAN", "TOO_RISKY"}
        assert item["score_breakdown"], "打分必须可追溯"
        assert item["reasons"], "必须给出人话解释"
        # 契约铁律 3：出现的数字都有来源（院校/专业/单位来源齐全）
        assert item["college"]["source_url"]
        assert item["major"]["source_url"]
        assert item["unit"]["unit_id"]

    stats = body["data"]["stats"]
    assert stats["hard_filtered_out"] >= 0
    assert 0.0 <= stats["data_coverage"] <= 1.0
    assert stats["rule"]["source_url"]
    assert body["evidence"], "整包响应必须带来源证据"


def test_recommend_excludes_too_risky_by_default(client: TestClient, student_id: str) -> None:
    default = client.post(f"{API}/recommend", json={"student_id": student_id, "limit": 50}).json()
    risky = client.post(
        f"{API}/recommend",
        json={"student_id": student_id, "limit": 50, "include_too_risky": True},
    ).json()
    default_tiers = {item["tier"] for item in default["data"]["items"]}
    assert "TOO_RISKY" not in default_tiers
    assert len(risky["data"]["items"]) >= len(default["data"]["items"])


# ---------------------------------------------------------------------------
# 志愿表：生成 → 读取 → 手改 → 校验 → 导出
# ---------------------------------------------------------------------------
@pytest.fixture()
def plan_id(client: TestClient, student_id: str) -> str:
    plan_id = f"plan-{uuid.uuid4().hex[:10]}"  # 唯一，避免重复运行撞主键
    response = client.post(
        f"{API}/plans/generate",
        json={"student_id": student_id, "plan_id": plan_id, "filters": {"regions": ["zhejiang"]}},
    )
    assert response.status_code == 200, response.text
    return plan_id


def test_plan_generate_and_get(client: TestClient, student_id: str, plan_id: str) -> None:
    explicit = f"plan-explicit-{uuid.uuid4().hex[:8]}"
    generated = client.post(
        f"{API}/plans/generate", json={"student_id": student_id, "plan_id": explicit}
    ).json()
    plan = generated["data"]["plan"]
    assert plan["id"] == explicit
    assert plan["items"], "平行志愿批次应生成志愿项"
    assert sum(plan["tier_distribution"].values()) == len(plan["items"])
    assert plan["rule"]["source_url"]
    assert generated["evidence"], "志愿表必须带逐项来源"
    # 每个志愿都必须能追溯到历史来源（或显式标注为新专业类比）
    unit_evidence = [entry for entry in generated["evidence"] if entry["what"] == "unit_history"]
    assert len(unit_evidence) >= len(plan["items"])
    assert all(entry.get("source_url") or entry.get("note") for entry in unit_evidence)

    fetched = client.get(f"{API}/plans/{plan_id}")
    assert fetched.status_code == 200
    assert fetched.json()["data"]["plan"]["id"] == plan_id
    assert client.get(f"{API}/plans/not-exist").status_code == 404


def test_plan_payload_carries_college_names(client: TestClient, plan_id: str) -> None:
    """志愿表的 PlanItem 只有 college_id；没有院校名，志愿表页面无法阅读（§8.2）。"""
    payload = client.get(f"{API}/plans/{plan_id}").json()["data"]
    colleges = payload["colleges"]
    assert colleges, "志愿表必须带院校索引"
    for item in payload["plan"]["items"]:
        college_id = item["unit"]["college_id"]
        assert college_id in colleges, college_id
        entry = colleges[college_id]
        assert entry["name"], college_id
        assert "source_url" in entry
    # 重新生成也必须带（generate / load / patch 三条路径口径一致）
    regenerated = client.post(
        f"{API}/plans/generate",
        json={"student_id": payload["plan"]["student_id"]},
    ).json()["data"]
    assert regenerated["colleges"], "生成时也必须带院校索引"


def test_plan_ids_are_unique_without_explicit_id(client: TestClient, student_id: str) -> None:
    """回归：不传 plan_id 时连续生成两张志愿表必须都成功（秒级时间戳 id 会同秒撞主键）。"""
    first = client.post(f"{API}/plans/generate", json={"student_id": student_id})
    second = client.post(f"{API}/plans/generate", json={"student_id": student_id})
    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    assert first.json()["data"]["plan"]["id"] != second.json()["data"]["plan"]["id"]


def test_plan_patch_reorder_and_validate(client: TestClient, plan_id: str) -> None:
    plan = client.get(f"{API}/plans/{plan_id}").json()["data"]["plan"]
    items = plan["items"]
    reordered = [items[1], items[0], *items[2:]]
    payload = {
        "items": [
            {"unit_id": item["unit"]["unit_id"], "obey_adjustment": item["obey_adjustment"]}
            for item in reordered
        ]
    }
    patched = client.patch(f"{API}/plans/{plan_id}/items", json=payload)
    assert patched.status_code == 200
    new_plan = patched.json()["data"]["plan"]
    assert new_plan["items"][0]["unit"]["unit_id"] == reordered[0]["unit"]["unit_id"]
    assert [item["position"] for item in new_plan["items"]] == list(range(1, len(new_plan["items"]) + 1))

    validated = client.post(f"{API}/plans/{plan_id}/validate")
    assert validated.status_code == 200
    assert "risks" in validated.json()["data"]

    unknown = client.patch(
        f"{API}/plans/{plan_id}/items",
        json={"items": [{"unit_id": "zhejiang-1900-NA-999999"}]},
    )
    assert unknown.status_code == 422
    assert unknown.json()["error"]["code"] == "UNKNOWN_UNITS"


def test_plan_remove_is_reversible(client: TestClient, plan_id: str) -> None:
    """回归：**删掉的志愿必须能加回来**。

    志愿填报工具里"移除"不可逆是危险缺陷——考生误删一个保底志愿却无法恢复。
    候选池 = 生成时的候选池（由同一套 evaluate_candidates 现场重算），因此：
    池内单位随时可加回，池外单位（含被硬约束剔除的）依然 422。
    """
    original = client.get(f"{API}/plans/{plan_id}").json()["data"]["plan"]
    removed = original["items"][0]
    removed_unit_id = removed["unit"]["unit_id"]
    rest = original["items"][1:]

    shrunk = client.patch(
        f"{API}/plans/{plan_id}/items",
        json={"items": [{"unit_id": item["unit"]["unit_id"]} for item in rest]},
    )
    assert shrunk.status_code == 200, shrunk.text
    remaining = [item["unit"]["unit_id"] for item in shrunk.json()["data"]["plan"]["items"]]
    assert removed_unit_id not in remaining

    # 加回来：追加到末尾，必须成功，且分层/概率/区间由重算结果给出（不沿用旧值）
    restored_items = [*remaining, removed_unit_id]
    restored = client.patch(
        f"{API}/plans/{plan_id}/items",
        json={"items": [{"unit_id": unit_id} for unit_id in restored_items]},
    )
    assert restored.status_code == 200, restored.text
    items = restored.json()["data"]["plan"]["items"]
    assert [item["unit"]["unit_id"] for item in items] == restored_items
    last = items[-1]
    assert last["tier"] in {"CHONG", "WEN", "BAO", "DIAN"}
    assert last["probability_interval"], "重算后必须带 ±1σ 概率区间（UI 只显示区间）"
    assert last["probability"] is not None


def test_plan_export_pdf_and_xlsx(client: TestClient, plan_id: str) -> None:
    pdf = client.get(f"{API}/plans/{plan_id}/export", params={"format": "pdf"})
    assert pdf.status_code == 200, pdf.text
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF"), "必须是真正的 PDF"
    assert len(pdf.content) > 1500

    xlsx = client.get(f"{API}/plans/{plan_id}/export", params={"format": "xlsx"})
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.content[:2] == b"PK", "xlsx 本质是 zip 包"
    assert len(xlsx.content) > 1500

    bad = client.get(f"{API}/plans/{plan_id}/export", params={"format": "docx"})
    assert bad.status_code == 422  # 只允许 pdf|xlsx


def test_plan_export_contains_disclaimer_and_sources(plan_id: str) -> None:
    """报告必须含免责声明与来源清单（§8 / §12）。"""
    import io

    from openpyxl import load_workbook

    from app.db import models as db
    from app.db import repositories as repo
    from app.services import plan_service, report_service

    with SessionLocal() as session:
        bundle = plan_service.load(session, plan_id)
        row = session.get(db.Student, bundle.plan.student_id)
        assert row is not None
        student = repo.row_to_student(row)
        report = report_service._with_college_names(
            session, report_service.build_report(session, bundle, student)
        )
    assert "严禁用于真实填报" in report["disclaimer"]
    assert report["sources"], "必须有来源清单"
    assert report["risks"] is not None

    workbook = load_workbook(io.BytesIO(report_service.render_xlsx(report)))
    assert {"志愿表", "历史证据", "风险提示", "说明与来源"} <= set(workbook.sheetnames)
    note_sheet = workbook["说明与来源"]
    text = "\n".join(str(cell.value) for row in note_sheet.iter_rows() for cell in row if cell.value)
    assert "免责声明" in text and "数据来源清单" in text


# ---------------------------------------------------------------------------
# 风险速查
# ---------------------------------------------------------------------------
def test_risk_scan(client: TestClient, student_id: str) -> None:
    plan = client.post(
        f"{API}/plans/generate", json={"student_id": student_id, "plan_id": f"plan-risk-{uuid.uuid4().hex[:6]}"}
    ).json()["data"]["plan"]
    unit_ids = [item["unit"]["unit_id"] for item in plan["items"][:5]]

    response = client.post(
        f"{API}/risk/scan",
        json={
            "student_id": student_id,
            "unit_ids": unit_ids,
            "obey_adjustment": True,
        },
    )
    assert response.status_code == 200
    body = response.json()
    assert body["data"]["scanned"] == len(unit_ids)
    for risk in body["data"]["risks"]:
        assert risk["code"] and risk["level"] in {"HIGH", "MEDIUM", "LOW"}
        assert risk["message"] and risk["suggestion"], "每个风险必须给可执行建议"
    assert body["evidence"], "风险扫描必须带证据链"

    unknown = client.post(
        f"{API}/risk/scan", json={"student_id": student_id, "unit_ids": ["zhejiang-1900-NA-999999"]}
    )
    assert unknown.status_code == 422


# ---------------------------------------------------------------------------
# 对话（SSE）
# ---------------------------------------------------------------------------
def test_chat_stream_and_history(client: TestClient) -> None:
    chat_service.reset()
    session_id = f"chat-{uuid.uuid4().hex[:8]}"
    response = client.post(
        f"{API}/chat", json={"message": "帮我看看志愿", "session_id": session_id, "student_id": "none"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/event-stream")
    text = response.text
    assert "event: start" in text and "event: delta" in text and "event: done" in text

    # 只检查正文（delta 拼接），避免把 session_id 里的数字当成模型输出
    deltas: list[str] = []
    for frame in text.split("\n\n"):
        if frame.startswith("event: delta"):
            payload = json.loads(frame.split("data: ", 1)[1])
            deltas.append(payload["text"])
    content = "".join(deltas)
    assert "以各省考试院官方文件与招生章程为准" in content  # 首次回复含免责声明（§12）
    # §0/§3.3：不得出现分数线/位次/录取率这类数字断言（"3+3"这类说明性文字不算）
    assert not re.search(r"\d{3,}", content), f"回复中不得出现三位以上数字：{content}"
    assert "%" not in content, f"回复中不得出现百分比：{content}"
    assert "宁可不答，不可编造" in content

    history = client.get(f"{API}/chat/{session_id}/history").json()
    messages = history["data"]["messages"]
    assert [message["role"] for message in messages] == ["user", "assistant"]
    assert messages[1]["missing_fields"], "缺档案时必须给出追问字段"

    empty = client.get(f"{API}/chat/chat-unknown/history").json()
    assert empty["data"]["count"] == 0 and empty["warnings"]


# ---------------------------------------------------------------------------
# 回测报告
# ---------------------------------------------------------------------------
def test_backtest_report_endpoint(client: TestClient) -> None:
    report_file = REPO_ROOT / "backtest_report.json"
    if not report_file.exists():
        pytest.skip("尚无 backtest_report.json：请先跑 scripts/run_backtest.py（离线任务）")
    payload = json.loads(report_file.read_text(encoding="utf-8"))

    body = client.get(
        f"{API}/backtest/report", params={"province": payload["province"], "year": payload["year"]}
    ).json()
    assert body["data"]["metrics"]["brier"] is not None
    assert set(body["data"]["checks"]) == {
        "safety_failure_rate",
        "wen_hit_rate",
        "chong_hit_rate",
        "brier",
    }
    assert body["evidence"]

    mismatch = client.get(f"{API}/backtest/report", params={"province": "hainan"})
    assert mismatch.status_code == 404
    assert mismatch.json()["error"]["code"] == "REPORT_UNAVAILABLE"
