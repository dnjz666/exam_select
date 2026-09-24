"""报告导出（L4，AGENTS.md §7 ``GET /plans/{id}/export?format=pdf|xlsx``）。

§8 / §12 强制要求：报告必须包含**档案 + 志愿表 + 每志愿依据 + 风险提示 + 免责声明 + 来源清单**。
xlsx 用 openpyxl；PDF 用 reportlab 的内置 CJK 字体 ``STSong-Light``（无需字体文件即可出中文）。
"""

from __future__ import annotations

import io
from html import escape
from collections.abc import Sequence

from sqlalchemy.orm import Session

from app.core.models import StudentProfile
from app.db import repositories as repo
from app.services import plan_service

DISCLAIMER = (
    "本报告由志愿填报智能体自动生成，仅供决策参考。最终请以各省教育考试院官方文件、"
    "招生计划与高校招生章程为准。请查看每项数据的来源标注；模拟数据仅用于流程演示，不可用于真实填报。"
)

TAG_NOTE = "每项计划与历史记录均标注数据性质和来源；涉及报考时请以考试院官方文件和招生章程为准。"

DATA_QUALITY_LABELS = {
    "OK": "记录完整",
    "DERIVED": "分数反查",
    "MISSING_RANK": "缺少位次",
    "COLLECTED": "征集志愿",
    "SUSPECT": "数据存疑",
}
VERIFIED_STATUS_LABELS = {
    "PRIMARY": "考试院原文已核实",
    "PRIMARY_GOV": "政府门户转述已核实",
    "SECONDARY": "转载来源待核实",
    "UNVERIFIED": "尚未核实",
}
RISK_LEVEL_LABELS = {"HIGH": "高风险", "MEDIUM": "中风险", "LOW": "提示"}
UNIT_TYPE_LABELS = {"MAJOR_COLLEGE": "专业+院校", "MAJOR_GROUP": "院校专业组"}


def _quality_label(value: str | None) -> str:
    return DATA_QUALITY_LABELS.get(value or "", "未标注")


def _label(mapping: dict[str, str], value: str | None) -> str:
    return mapping.get(value or "", "未标注")


def _source_note(source_flags: list[bool]) -> str:
    if not source_flags:
        return "当前报告未包含志愿或历史数据；请先生成志愿表。"
    if any(source_flags) and any(not flag for flag in source_flags):
        return "当前报告同时包含真实来源数据与模拟数据；各条数据性质请见对应记录。"
    if any(source_flags):
        return "当前报告包含模拟数据；请勿用于真实填报。"
    return "当前计划与历史记录均标记为真实来源；仍须核对考试院官方文件和招生章程。"


def build_report(session: Session, bundle: plan_service.PlanBundle, student: StudentProfile) -> dict:
    """组装结构化报告（xlsx / PDF / 前端报告页共用同一份数据）。"""
    plan = bundle.plan
    evidence_by_unit: dict[str, list[dict]] = {}
    for entry in bundle.evidence:
        if entry.get("what") == "unit_history":
            evidence_by_unit.setdefault(entry["unit_id"], []).append(entry)

    items: list[dict] = []
    for item in plan.items:
        unit = item.unit
        interval = None
        if item.probability is not None:
            # 报告统一以区间呈现（§8 UI 强制要求）
            low = max(0.02, item.probability - 0.05)
            high = min(0.98, item.probability + 0.05)
            interval = [round(low, 4), round(high, 4)]
        items.append(
            {
                "position": item.position,
                "college_id": unit.college_id,
                "unit_id": unit.unit_id,
                "college_name": "",  # 由 _with_college_names 填充（院校名/城市/层次均来自 colleges 表）
                "major_name": unit.major_name,
                "group_name": unit.group_name,
                "unit_type": unit.unit_type.value,
                "batch": unit.batch,
                "plan_count": unit.plan_count,
                "tuition": unit.tuition,
                "source_url": unit.source_url,
                "is_synthetic": unit.is_synthetic,
                "tier": item.tier.value,
                "probability": item.probability,
                "probability_interval": interval,
                "utility": item.utility,
                "obey_adjustment": item.obey_adjustment,
                "reasons": list(item.notes),  # 概率模型给的人话理由（§7"每志愿依据"）
                "evidence": evidence_by_unit.get(unit.unit_id, []),
            }
        )

    sources = sorted(
        {entry["source_url"] for entry in bundle.evidence if entry.get("source_url")}
        | {item["source_url"] for item in items if item.get("source_url")}
        | {plan.rule.source_url}
    )
    source_flags = [item["is_synthetic"] for item in items]
    source_flags.extend(bool(entry.get("is_synthetic", True)) for item in items for entry in item["evidence"])
    report = {
        "generated_at": repo.now_iso(),
        "student": {
            "id": student.id,
            "province": student.province,
            "year": student.year,
            "subjects": list(student.subjects),
            "total_score": student.total_score,
            "rank": student.rank,
            "foreign_language": student.foreign_language,
        },
        "rule": {
            "province": plan.rule.province,
            "batch_code": plan.rule.batch_code,
            "batch_name": plan.rule.batch_name,
            "max_volunteers": plan.rule.max_volunteers,
            "majors_per_group": plan.rule.majors_per_group,
            "has_major_adjustment": plan.rule.has_major_adjustment,
            "is_parallel": plan.rule.is_parallel,
            "verified_status": plan.rule.verified_status.value,
            "verified_year": plan.rule.verified_year,
            "source_url": plan.rule.source_url,
        },
        "plan": {
            "id": plan.id,
            "items": items,
            "tier_distribution": plan.tier_distribution,
            "total_utility": plan.total_utility,
            "violations": [violation.model_dump(mode="json") for violation in plan.violations],
        },
        "risks": [risk.model_dump(mode="json") for risk in bundle.risks],
        "warnings": list(bundle.warnings),
        "sources": sources,
        "source_note": _source_note(source_flags),
        "data_sources": {
            "contains_synthetic": any(source_flags),
            "contains_real": any(not flag for flag in source_flags),
            "scope": "志愿表当前计划与所用历史证据",
        },
        "disclaimer": DISCLAIMER,
    }
    return report


def _with_college_names(session: Session, report: dict) -> dict:
    colleges = repo.load_colleges(session)
    majors = repo.load_majors(session)
    for item in report["plan"]["items"]:
        college = colleges.get(item["college_id"])
        item["college_name"] = college.name if college else item["college_id"]
        item["college_city"] = college.city if college else ""
        item["level_tags"] = "、".join(college.level_tags) if college else ""
        item["college_source_url"] = college.source_url if college else ""
        item["tuition_note"] = "民办/中外合作，请在卡片核验学费" if college and not college.is_public else ""
        if college and college.source_url:
            report["sources"].append(college.source_url)
    report["sources"] = sorted(set(report["sources"]))
    _ = majors
    return report


def render_xlsx(report: dict) -> bytes:
    """生成 xlsx（四张表：志愿表 / 历史证据 / 风险 / 说明与来源）。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "志愿表"
    header = ["顺序", "院校", "城市", "层次", "专业", "专业组", "投档单位", "计划数", "学费(元/年)", "计划来源", "数据性质",
              "分层", "概率", "概率区间", "服从调剂"]
    sheet.append(header)
    for item in report["plan"]["items"]:
        interval = item.get("probability_interval")
        sheet.append(
            [
                item["position"],
                item["college_name"],
                item.get("college_city", ""),
                item.get("level_tags", ""),
                item["major_name"],
                item.get("group_name") or "—",
                _label(UNIT_TYPE_LABELS, item["unit_type"]),
                item["plan_count"],
                item["tuition"],
                item.get("source_url", ""),
                "模拟数据" if item.get("is_synthetic", True) else "真实来源",
                item["tier"],
                item["probability"],
                f"{interval[0]:.0%}–{interval[1]:.0%}" if interval else "—",
                "—" if item["obey_adjustment"] is None else ("是" if item["obey_adjustment"] else "否"),
            ]
        )

    evidence_sheet = workbook.create_sheet("历史证据")
    evidence_sheet.append(["志愿顺序", "投档单位", "年份", "最低位次", "数据质量", "数据性质", "来源", "备注"])
    for item in report["plan"]["items"]:
        if not item["evidence"]:
            evidence_sheet.append(
                [item["position"], item["unit_id"], "—", "—", "无本单位历史", "—", "—",
                 "无本单位历史，概率参考同类单位；请在推荐依据中核对类比证据"]
            )
        for entry in item["evidence"]:
            evidence_sheet.append(
                [
                    item["position"],
                    entry["unit_id"],
                    entry.get("year", ""),
                    entry.get("min_rank", ""),
                    _quality_label(entry.get("data_quality")),
                    "模拟数据" if entry.get("is_synthetic", True) else "真实来源",
                    entry.get("source_url", ""),
                    entry.get("note", ""),
                ]
            )

    risk_sheet = workbook.create_sheet("风险提示")
    risk_sheet.append(["等级", "投档单位", "说明", "建议"])
    for risk in report["risks"]:
        risk_sheet.append(
            [_label(RISK_LEVEL_LABELS, risk["level"]), risk.get("unit_id") or "—", risk["message"], risk["suggestion"]]
        )

    note_sheet = workbook.create_sheet("说明与来源")
    note_sheet.append(["生成时间", report["generated_at"]])
    note_sheet.append(["省份/批次", f"{report['rule']['province']} / {report['rule']['batch_name']}"])
    note_sheet.append(["规则核实状态", _label(VERIFIED_STATUS_LABELS, report["rule"]["verified_status"])])
    note_sheet.append(["考生位次", report["student"]["rank"]])
    note_sheet.append([])
    note_sheet.append(["数据来源清单"])
    note_sheet.append([report["source_note"] + " " + TAG_NOTE])
    for url in report["sources"]:
        note_sheet.append([url])
    note_sheet.append([])
    note_sheet.append(["免责声明", DISCLAIMER])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(report: dict) -> bytes:
    """生成中文 PDF（reportlab 内置 STSong-Light CID 字体）。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    title_style = ParagraphStyle("cn-title", fontName="STSong-Light", fontSize=16, leading=22)
    body_style = ParagraphStyle("cn-body", fontName="STSong-Light", fontSize=9.5, leading=14)
    small_style = ParagraphStyle("cn-small", fontName="STSong-Light", fontSize=8, leading=12,
                                 textColor=colors.HexColor("#555555"))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="高考志愿表报告",
    )
    story: list = [Paragraph("高考志愿填报报告", title_style), Spacer(1, 6)]
    student = report["student"]
    rule = report["rule"]
    story.append(
        Paragraph(
            f"考生：{student['id']}　省份：{student['province']}　年份：{student['year']}　"
            f"选考：{'、'.join(student['subjects'])}　总分：{student['total_score']}　位次：{student['rank']}",
            body_style,
        )
    )
    story.append(
        Paragraph(
            f"批次：{rule['batch_name']}（{rule['batch_code']}）　平行志愿数：{rule['max_volunteers']}　"
            f"规则核实状态：{_label(VERIFIED_STATUS_LABELS, rule['verified_status'])}　来源：{rule['source_url']}",
            small_style,
        )
    )
    if rule["verified_status"] in {"SECONDARY", "UNVERIFIED"}:
        story.append(
            Paragraph(
                "<b>⚠️ 规则待核实</b>：该省规则尚未达到考试院官方原文核实等级，"
                "本报告不得用于真实填报。",
                body_style,
            )
        )
    story.append(Spacer(1, 8))

    rows = [["#", "院校", "专业", "分层", "概率区间", "计划", "学费/年", "数据性质", "计划来源", "服从调剂"]]
    for item in report["plan"]["items"]:
        interval = item.get("probability_interval")
        plan_source = item.get("source_url") or ""
        if plan_source.startswith(("http://", "https://")):
            plan_source_cell = Paragraph(f'<link href="{escape(plan_source, quote=True)}">查看来源</link>', small_style)
        elif plan_source:
            plan_source_cell = Paragraph("模拟来源标识", small_style)
        else:
            plan_source_cell = "未提供来源"
        rows.append(
            [
                str(item["position"]),
                item["college_name"],
                item["major_name"],
                item["tier"],
                f"{interval[0]:.0%}–{interval[1]:.0%}" if interval else "—",
                str(item["plan_count"]),
                (
                    f"{item['tuition']:,} 元/年"
                    if isinstance(item.get("tuition"), int) and item["tuition"] > 0
                    else "未收录（请核对招生章程）"
                ),
                "模拟" if item.get("is_synthetic", True) else "真实",
                plan_source_cell,
                "—" if item["obey_adjustment"] is None else ("是" if item["obey_adjustment"] else "否"),
            ]
        )
    table = Table(rows, colWidths=[8 * mm, 23 * mm, 23 * mm, 12 * mm, 20 * mm, 9 * mm, 27 * mm, 18 * mm, 18 * mm, 14 * mm])
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story += [table, Spacer(1, 10)]

    history_rows = [["#", "院校 / 专业", "依据类型", "年份", "最低位次", "数据质量", "数据性质", "来源"]]
    for item in report["plan"]["items"]:
        unit_label = f"{item['college_name']} · {item['major_name']}"
        if not item.get("evidence"):
            history_rows.append(
                [str(item["position"]), Paragraph(escape(unit_label), small_style), "无本单位历史",
                 "—", "—", "—", "—", "概率参考同类单位；请在推荐依据中核对类比证据"]
            )
        for entry in item.get("evidence", []):
            source_url = entry.get("source_url") or ""
            if source_url.startswith(("http://", "https://")):
                source_cell = Paragraph(f'<link href="{escape(source_url, quote=True)}">查看来源</link>', small_style)
            elif source_url:
                # synthetic:// 等内部来源标识也要明确显示，不能伪装成缺来源。
                source_cell = Paragraph(escape(source_url), small_style)
            else:
                source_cell = "未提供来源"
            history_rows.append(
                [
                    str(item["position"]),
                    Paragraph(escape(unit_label), small_style),
                    "同类单位类比" if entry.get("note") else "本单位历史",
                    str(entry.get("year") or "—"),
                    f"{entry['min_rank']:,}" if isinstance(entry.get("min_rank"), int) else "—",
                    _quality_label(entry.get("data_quality")),
                    "模拟数据" if entry.get("is_synthetic", True) else "真实来源",
                    source_cell,
                ]
            )
    if len(history_rows) > 1:
        story.append(Paragraph("历史证据", body_style))
        history_table = Table(
            history_rows,
            colWidths=[8 * mm, 34 * mm, 24 * mm, 12 * mm, 18 * mm, 20 * mm, 20 * mm, 30 * mm],
            repeatRows=1,
        )
        history_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        story += [history_table, Spacer(1, 10)]

    if report["risks"]:
        story.append(Paragraph("风险提示", body_style))
        risk_rows = [["等级", "说明", "建议"]]
        for risk in report["risks"]:
            risk_rows.append([_label(RISK_LEVEL_LABELS, risk["level"]), risk["message"], risk["suggestion"]])
        risk_table = Table(risk_rows, colWidths=[18 * mm, 80 * mm, 80 * mm])
        risk_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ]
            )
        )
        story += [risk_table, Spacer(1, 10)]

    story.append(Paragraph("数据来源清单", body_style))
    for url in report["sources"]:
        story.append(Paragraph(f"· {url}", small_style))
    story.append(Paragraph(report["source_note"], small_style))
    story += [Spacer(1, 8), Paragraph(f"免责声明：{DISCLAIMER}", small_style)]

    doc.build(story)
    return buffer.getvalue()


def export(session: Session, plan_id: str, fmt: str) -> tuple[bytes, str, str]:
    """导出报告：返回 ``(内容, media_type, 文件名)``。"""
    bundle = plan_service.load(session, plan_id)
    row = repo.get_student(session, bundle.plan.student_id)
    if row is None:
        raise plan_service.PlanNotFound(f"student:{bundle.plan.student_id}")
    student = repo.row_to_student(row)
    report = _with_college_names(session, build_report(session, bundle, student))

    if fmt == "xlsx":
        return (
            render_xlsx(report),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"volunteer-plan-{plan_id}.xlsx",
        )
    if fmt == "pdf":
        return render_pdf(report), "application/pdf", f"volunteer-plan-{plan_id}.pdf"
    raise ValueError(f"不支持的报告格式：{fmt}")


__all__ = ["DISCLAIMER", "build_report", "export", "render_pdf", "render_xlsx"]

_ = Sequence
